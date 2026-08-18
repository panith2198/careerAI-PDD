from __future__ import annotations

import ast
import json
import os
import re
from collections import Counter
from dataclasses import dataclass, asdict
from datetime import datetime
from pathlib import Path
from typing import Iterable
from urllib import request, error

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
BACKEND = ROOT / "backend"
OUT = ROOT / "Vulnerability Test Results"


@dataclass
class Endpoint:
    endpoint: str
    method: str
    auth_required: str
    expected_roles: str
    controller_file_path: str


@dataclass
class Finding:
    severity: str
    vulnerability_type: str
    file_path: str
    endpoint: str
    description: str
    exploitation_scenario: str
    impact: str
    recommended_fix: str


def parse_router_prefixes() -> dict[str, list[str]]:
    main_py = BACKEND / "main.py"
    source = main_py.read_text(encoding="utf-8")
    aliases: dict[str, str] = {}
    prefixes: dict[str, list[str]] = {}

    for match in re.finditer(r"from app\.api\.v1\.(\w+) import router as (\w+)", source):
        aliases[match.group(2)] = match.group(1)

    for match in re.finditer(r"app\.include_router\((\w+), prefix=\"([^\"]+)\"", source):
        module = aliases.get(match.group(1))
        if module:
            prefixes.setdefault(module, []).append(match.group(2))

    return prefixes


def decorator_path(call: ast.Call) -> tuple[str, str] | None:
    if not isinstance(call.func, ast.Attribute):
        return None
    method = call.func.attr.upper()
    if method not in {"GET", "POST", "PUT", "PATCH", "DELETE", "WEBSOCKET"}:
        return None
    if not call.args:
        return None
    first = call.args[0]
    if isinstance(first, ast.Constant) and isinstance(first.value, str):
        return method, first.value
    return None


def function_auth_required(func: ast.AST) -> str:
    text = ast.unparse(func) if hasattr(ast, "unparse") else ""
    auth_markers = ["get_current_user", "oauth2_scheme", "current_user"]
    return "Yes" if any(marker in text for marker in auth_markers) else "No/Unknown"


def discover_endpoints() -> list[Endpoint]:
    prefixes = parse_router_prefixes()
    endpoints: list[Endpoint] = [
        Endpoint("/health", "GET", "No", "Public", "backend/main.py"),
    ]

    for file in sorted((BACKEND / "app" / "api" / "v1").glob("*.py")):
        module = file.stem
        if module in {"__init__", "deps"}:
            continue
        source = file.read_text(encoding="utf-8")
        tree = ast.parse(source)
        module_prefixes = prefixes.get(module, ["/api/v1"])
        for node in ast.walk(tree):
            if not isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
                continue
            for deco in node.decorator_list:
                if isinstance(deco, ast.Call):
                    parsed = decorator_path(deco)
                    if parsed:
                        method, route = parsed
                        for prefix in module_prefixes:
                            full = (prefix.rstrip("/") + "/" + route.lstrip("/")).rstrip("/")
                            endpoints.append(
                                Endpoint(
                                    endpoint=full or "/",
                                    method=method,
                                    auth_required=function_auth_required(node),
                                    expected_roles="student/user" if function_auth_required(node) == "Yes" else "Public/Unknown",
                                    controller_file_path=str(file.relative_to(ROOT)),
                                )
                            )

    return sorted(endpoints, key=lambda item: (item.endpoint, item.method))


def has_text(path: Path, pattern: str) -> bool:
    return bool(re.search(pattern, path.read_text(encoding="utf-8", errors="ignore"), re.MULTILINE))


def static_findings() -> list[Finding]:
    findings: list[Finding] = []
    main_py = BACKEND / "main.py"
    config_py = BACKEND / "app" / "core" / "config.py"
    deps_py = BACKEND / "app" / "api" / "v1" / "deps.py"
    auth_py = BACKEND / "app" / "api" / "v1" / "auth.py"

    if has_text(main_py, r"allow_origins=\[\"\*\"\].*allow_credentials=True|allow_credentials=True"):
        findings.append(Finding(
            "Medium",
            "CORS Misconfiguration",
            "backend/main.py",
            "All API endpoints",
            "CORS is configured broadly and allows credentials. Browsers may reject wildcard credentials, and overly broad origins increase cross-site attack surface.",
            "A malicious web origin can attempt credentialed API requests if future browser or proxy behavior changes or if origins are relaxed incorrectly.",
            "Increases account and token exposure risk for browser clients.",
            "Configure an explicit allowlist such as https://careerai.welcos.in and local development origins only.",
        ))

    if has_text(config_py, r"SECRET_KEY.*secretkeyplaceholder"):
        findings.append(Finding(
            "High",
            "Insecure Secret Default",
            "backend/app/core/config.py",
            "JWT authentication",
            "The application ships with a weak default JWT signing secret.",
            "If production is deployed without overriding SECRET_KEY, attackers can forge tokens.",
            "Full account impersonation and privilege escalation.",
            "Require SECRET_KEY through environment variables in production and fail startup when it is unset or placeholder.",
        ))

    if has_text(config_py, r"DEBUG: bool = Field\(default=True\)"):
        findings.append(Finding(
            "Medium",
            "Unsafe Production Default",
            "backend/app/core/config.py",
            "Application startup",
            "DEBUG defaults to true.",
            "A missed environment override may expose verbose errors and internals.",
            "Information disclosure and easier exploit development.",
            "Default DEBUG to false and enable only in local development.",
        ))

    if has_text(auth_py, r"jwt\.decode\(payload\.refresh_token.*verify_exp.*False"):
        findings.append(Finding(
            "High",
            "JWT Expiration Bypass",
            "backend/app/api/v1/auth.py",
            "/api/v1/auth/refresh",
            "Refresh token logic decodes JWTs with expiration verification disabled and falls back to user id 1 on errors.",
            "An attacker can replay expired tokens or force fallback behavior to mint a token for a predictable user.",
            "Account takeover risk, especially for the first user.",
            "Implement real refresh tokens, verify expiration/signature, store rotation state, and never fall back to user id 1.",
        ))

    if has_text(deps_py, r"print\(f\"\[DEBUG AUTH\].*Token"):
        findings.append(Finding(
            "High",
            "Token Leakage In Logs",
            "backend/app/api/v1/deps.py",
            "Authenticated APIs",
            "Authentication dependency prints JWT token material to logs.",
            "Anyone with log access can recover bearer token prefixes or potentially full tokens if logging is expanded.",
            "Session hijacking and sensitive data exposure.",
            "Remove token logging and use structured security-safe debug events.",
        ))

    if not has_text(main_py, r"CoreMiddleware|RateLimitMiddleware"):
        findings.append(Finding(
            "Medium",
            "Rate Limiting Not Registered",
            "backend/main.py",
            "Authentication and API endpoints",
            "Rate limiting middleware exists but is not registered in the FastAPI app.",
            "Attackers can brute-force login or flood expensive AI endpoints.",
            "Credential stuffing and denial of wallet/service risk.",
            "Register rate limiting middleware and add endpoint-specific limits for login, OTP, file upload, and AI endpoints.",
        ))

    if not has_text(main_py, r"Strict-Transport-Security|Content-Security-Policy|X-Content-Type-Options"):
        findings.append(Finding(
            "Low",
            "Missing Security Headers",
            "backend/main.py",
            "All API endpoints",
            "The backend does not add standard hardening headers.",
            "Clients receive fewer browser-side protections against MIME sniffing and framing.",
            "Defense-in-depth gap.",
            "Add security headers at the proxy and/or FastAPI middleware.",
        ))

    return findings


def dependency_rows() -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for file in [BACKEND / "requirements.txt", ROOT / "web" / "package.json", ROOT / "mobile" / "package.json"]:
        if not file.exists():
            continue
        if file.suffix == ".txt":
            for line in file.read_text(encoding="utf-8").splitlines():
                line = line.strip()
                if line and not line.startswith("#"):
                    rows.append({"Source": str(file.relative_to(ROOT)), "Package": line, "Scanner Status": "Cataloged", "Finding": "Run CI scanners for CVE confirmation"})
        else:
            data = json.loads(file.read_text(encoding="utf-8"))
            for section in ["dependencies", "devDependencies"]:
                for name, version in data.get(section, {}).items():
                    rows.append({"Source": str(file.relative_to(ROOT)), "Package": f"{name}@{version}", "Scanner Status": "Cataloged", "Finding": "Run CI scanners for CVE confirmation"})
    return rows


def dast_rows(api_url: str) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    probes = [
        ("Health check", "GET", "/health"),
        ("CORS preflight register", "OPTIONS", "/api/v1/auth/register"),
        ("Unauthenticated protected endpoint", "GET", "/api/v1/users/me"),
    ]
    for name, method, path in probes:
        url = api_url.rstrip("/") + path
        req = request.Request(url, method=method)
        req.add_header("Origin", "http://127.0.0.1:5175")
        if method == "OPTIONS":
            req.add_header("Access-Control-Request-Method", "POST")
            req.add_header("Access-Control-Request-Headers", "content-type,authorization")
        try:
            with request.urlopen(req, timeout=10) as response:
                rows.append({"Probe": name, "Method": method, "URL": url, "Status": str(response.status), "Result": response.read(200).decode("utf-8", "ignore")})
        except error.HTTPError as exc:
            rows.append({"Probe": name, "Method": method, "URL": url, "Status": str(exc.code), "Result": exc.read(200).decode("utf-8", "ignore")})
        except Exception as exc:
            rows.append({"Probe": name, "Method": method, "URL": url, "Status": "ERROR", "Result": str(exc)})
    return rows


def write_markdown(endpoints: list[Endpoint], findings: list[Finding], deps: list[dict[str, str]], dast: list[dict[str, str]]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    counts = Counter(f.severity for f in findings)
    total = len(findings)
    score = max(0, 100 - counts["Critical"] * 25 - counts["High"] * 15 - counts["Medium"] * 7 - counts["Low"] * 3)

    inventory = f"""# Backend Inventory

- Framework: FastAPI
- Language: Python
- API architecture: REST APIs with FastAPI routers under `/api/v1`
- Authentication: JWT bearer tokens signed with HS256
- Authorization model: user-context dependency, limited explicit role checks observed
- Database technology: MySQL 8 compatible database
- ORM usage: SQLAlchemy async ORM with aiomysql
- API documentation: FastAPI OpenAPI at `/openapi.json`
- Middleware: CORS middleware; custom rate-limit middleware exists but is not registered
- File upload functionality: resume upload, avatar upload, and RAG knowledge-base upload endpoints
- Session handling: stateless JWT access tokens; refresh endpoint is implemented in auth router
- Third-party integrations: Mistral AI, Redis/Celery, SMTP, Razorpay
"""

    review = [inventory, "\n# Security Findings\n"]
    for idx, finding in enumerate(findings, 1):
        review.append(f"""## {idx}. {finding.vulnerability_type}

- Severity: {finding.severity}
- File Path: `{finding.file_path}`
- Endpoint: `{finding.endpoint}`
- Description: {finding.description}
- Exploitation Scenario: {finding.exploitation_scenario}
- Impact: {finding.impact}
- Recommended Fix: {finding.recommended_fix}
""")
    (OUT / "security-review.md").write_text("\n".join(review), encoding="utf-8")

    summary = f"""# Executive Summary

Total Findings: {total}

Critical: {counts['Critical']}
High: {counts['High']}
Medium: {counts['Medium']}
Low: {counts['Low']}

Most Critical Risks

1. JWT refresh token validation bypass and fallback behavior.
2. Insecure default secret and debug settings if production env variables are missed.
3. Token material can be written to logs by the authentication dependency.

Overall Security Score

{score}/100
"""
    (OUT / "executive-summary.md").write_text(summary, encoding="utf-8")

    dep_lines = ["# Dependency Report", "", "Local scanners available in this run: Python package catalog parser. CI workflow installs Semgrep, Trivy, and Gitleaks.", ""]
    dep_lines.extend(f"- `{row['Source']}`: {row['Package']} - {row['Finding']}" for row in deps[:250])
    dep_lines.append("\n## DAST Probe Results\n")
    dep_lines.extend(f"- {row['Probe']} `{row['Method']} {row['URL']}` -> {row['Status']}" for row in dast)
    (OUT / "dependency-report.md").write_text("\n".join(dep_lines), encoding="utf-8")


def autosize(sheet) -> None:
    for col_idx, column in enumerate(sheet.columns, 1):
        max_len = 0
        for cell in column:
            max_len = max(max_len, len(str(cell.value or "")))
        sheet.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 12), 70)


def write_workbook(path: Path, endpoints: list[Endpoint], findings: list[Finding], deps: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheets = [
        ("Security Findings", [asdict(f) for f in findings]),
        ("Endpoint Inventory", [asdict(e) for e in endpoints]),
        ("Dependency Vulnerabilities", deps),
        ("Risk Summary", [{"Severity": sev, "Count": count} for sev, count in Counter(f.severity for f in findings).items()]),
    ]
    workbook.remove(workbook.active)
    for name, rows in sheets:
        sheet = workbook.create_sheet(name)
        if rows:
            headers = list(rows[0].keys())
            sheet.append(headers)
            for row in rows:
                sheet.append([row.get(header, "") for header in headers])
        else:
            sheet.append(["No records"])
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="1F2937")
        autosize(sheet)
    workbook.save(path)


def main() -> None:
    endpoints = discover_endpoints()
    findings = static_findings()
    deps = dependency_rows()
    api_url = os.environ.get("SECURITY_API_URL", "https://careerai-api.welcos.in")
    dast = dast_rows(api_url)

    write_markdown(endpoints, findings, deps, dast)
    write_workbook(OUT / "findings.xlsx", endpoints, findings, deps)
    write_workbook(OUT / "endpoint-inventory.xlsx", endpoints, findings, deps)
    print(f"Generated security reports in {OUT}")
    print(f"Endpoints discovered: {len(endpoints)}")
    print(f"Findings discovered: {len(findings)}")


if __name__ == "__main__":
    main()
