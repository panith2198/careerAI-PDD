from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
findings_file = ROOT / "Vulnerability Test Results" / "findings.xlsx"

if not findings_file.exists():
    raise SystemExit("findings.xlsx was not generated")

workbook = load_workbook(findings_file)
sheet = workbook["Security Findings"]
headers = [cell.value for cell in sheet[1]]
severity_index = headers.index("severity") + 1
critical_count = 0

for row in range(2, sheet.max_row + 1):
    if sheet.cell(row=row, column=severity_index).value == "Critical":
        critical_count += 1

if critical_count:
    raise SystemExit(f"Critical vulnerabilities found: {critical_count}")

print("No Critical vulnerabilities found.")
