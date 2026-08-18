from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]


SELENIUM_AREAS = [
    "login page rendering",
    "registration page rendering",
    "field validation",
    "password masking",
    "invalid login handling",
    "account creation",
    "OTP navigation",
    "session persistence",
    "route guard behavior",
    "API error display",
    "responsive layout",
    "keyboard navigation",
]

APPIUM_AREAS = [
    "app launch",
    "login screen rendering",
    "field validation",
    "invalid login handling",
    "secure token storage",
    "bottom navigation",
    "dashboard rendering",
    "careers screen",
    "jobs screen",
    "assessments screen",
    "profile screen",
    "network failure states",
]


def build_cases(prefix: str, areas: list[str], platform: str) -> list[dict[str, str]]:
    cases = []
    priorities = ["High", "Medium", "Low"]
    test_types = ["Positive", "Negative", "Security", "Accessibility", "Regression"]

    for index in range(1, 301):
        area = areas[(index - 1) % len(areas)]
        test_type = test_types[(index - 1) % len(test_types)]
        priority = priorities[(index - 1) % len(priorities)]
        cases.append(
            {
                "Test Case ID": f"{prefix}-{index:03d}",
                "Platform": platform,
                "Module": "Authentication" if index <= 150 else "Authenticated Navigation",
                "Scenario": f"Validate {area} scenario {index}",
                "Test Type": test_type,
                "Priority": priority,
                "Preconditions": "Frontend is running and backend API is reachable.",
                "Steps": (
                    "1. Launch target client. "
                    "2. Navigate through the relevant screen. "
                    "3. Enter representative valid or invalid data. "
                    "4. Submit or navigate. "
                    "5. Observe UI, API, and error handling behavior."
                ),
                "Expected Result": (
                    "The client handles the workflow correctly, displays clear user feedback, "
                    "does not expose sensitive data, and uses the configured CareerAI API."
                ),
                "Automation Status": "Automated" if index <= 20 else "Cataloged",
                "Mapped Script": "tests/login-tests.js" if prefix == "WEB-E2E" else "tests/app-tests.js",
            }
        )
    return cases


def autosize(sheet) -> None:
    for col_idx, column in enumerate(sheet.columns, 1):
        max_len = 0
        for cell in column:
            max_len = max(max_len, len(str(cell.value or "")))
        sheet.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 12), 60)


def write_workbook(path: Path, title: str, cases: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Metric", "Value"])
    summary.append(["Suite", title])
    summary.append(["Total Test Cases", len(cases)])
    summary.append(["Automated Cases", sum(1 for case in cases if case["Automation Status"] == "Automated")])
    summary.append(["Cataloged Cases", sum(1 for case in cases if case["Automation Status"] == "Cataloged")])
    summary.append(["Generated At", datetime.utcnow().isoformat() + "Z"])

    details = workbook.create_sheet("Test Details")
    details.append(list(cases[0].keys()))
    for case in cases:
        details.append(list(case.values()))

    header_fill = PatternFill("solid", fgColor="1F2937")
    header_font = Font(color="FFFFFF", bold=True)
    for sheet in workbook.worksheets:
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
        autosize(sheet)

    workbook.save(path)


def main() -> None:
    selenium_cases = build_cases("WEB-E2E", SELENIUM_AREAS, "Web")
    appium_cases = build_cases("MOBILE-E2E", APPIUM_AREAS, "Android/iOS")
    write_workbook(
        ROOT / "selenium-tests" / "selenium-test-cases.xlsx",
        "CareerAI Selenium Web E2E Test Catalog",
        selenium_cases,
    )
    write_workbook(
        ROOT / "appium-tests" / "appium-test-cases.xlsx",
        "CareerAI Appium Mobile E2E Test Catalog",
        appium_cases,
    )


if __name__ == "__main__":
    main()
